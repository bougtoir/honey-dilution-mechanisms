"""Generate the supplementary material DOCX."""
import os
import re

import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "manuscript")

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(10)


def h(t, l=1):
    p = doc.add_heading(t, level=l)
    for r in p.runs:
        r.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0, 0, 0)
    return p


def add_df(df, cols=None, maxrows=200):
    if cols:
        df = df[cols]
    df = df.head(maxrows)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for j, c in enumerate(df.columns):
        t.rows[0].cells[j].text = str(c)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = "" if pd.isna(v) else str(v)[:80]
    doc.add_paragraph()


h("Supplementary Material", 0)
doc.add_paragraph("Dilution-Dependent Antibacterial Mechanisms of Honey: "
                  "A Systematic Review and Cross-Study Quantitative Evidence Synthesis")

h("S1. Search strategy", 1)
doc.add_paragraph(
    "Databases: PubMed (NCBI E-utilities esearch/efetch), Europe PMC REST API, "
    "Crossref works API. Queries combined honey terms (honey, manuka, "
    "leptospermum, melipona, stingless bee) with antibacterial endpoints (MIC, "
    "MBC, antibacterial, antimicrobial, inhibition zone) and mechanistic terms "
    "(hydrogen peroxide, glucose oxidase, methylglyoxal, water activity, pH, "
    "dilution, catalase). Exact query strings are stored verbatim in "
    "data/raw/search/ alongside raw JSON responses; script scripts/00_search.py "
    "reproduces them. Backward citation mining via Europe PMC "
    "(scripts/00b_citation_mining.py) appended 134 records.")
doc.add_paragraph(
    "Stage-1 exclusion categories applied to title/abstract: not about honey "
    "(n = 750), no antibacterial outcome (436), no abstract (103), "
    "review/non-primary (61), off-topic (21).")

h("S2. Eligibility rules", 1)
doc.add_paragraph(
    "Deterministic classifier (scripts/04_eligibility.py): records required a "
    "honey term AND an antibacterial-endpoint term AND at least one mechanistic "
    "or quantitative flag (dilution/concentration series, MIC/MBC value, zone "
    "measurement, catalase, peroxide, methylglyoxal, water activity, artificial/"
    "simulated honey comparator, physicochemical characterization). Records "
    "lacking a mechanistic flag but otherwise eligible were marked review_flag.")

h("S3. Extraction method detail", 1)
doc.add_paragraph(
    "PMC JATS XML tables were parsed with colspan/rowspan expansion and nested "
    "<alternatives> tables resolved (scripts/05_parse_tables.py). The "
    "rule-based extractor (scripts/06_autoextract.py) recognizes three dominant "
    "layouts: (A) honey-by-organism MIC/MBC/zone tables, including two-row "
    "headers with MIC/MBC subcolumns and transposed orientations; (B) "
    "honey-by-parameter physicochemical tables; (C) concentration-series tables "
    "in either row- or column-oriented form. Organism abbreviations were "
    "harmonized by a documented regex map (scripts/08_build_dataset.py). Key "
    "mechanistic tables (catalase pairs, simulated-honey comparators, MGO "
    "dose-response) were additionally curated manually after inspection of the "
    "source table and caption (scripts/07_manual_extraction.py); curated rows "
    "take precedence and are flagged extract_method=manual_curated.")

h("S4. Risk-of-bias framework", 1)
doc.add_paragraph(
    "Domains (0 = high concern, 1 = some, 2 = low): D1 assay standardization "
    "(recognized MIC/MBC method); D2 comparator quality (artificial honey, "
    "sugar-equivalent or vehicle control); D3 dispersion reporting (SD/SE/n in "
    "extracted table); D5 mechanistic perturbation (catalase, H2O2, MGO or "
    "water-activity measurement); D6 reporting completeness (MIC+MBC or zone "
    "with endpoint). Concentration-series availability is recorded separately. "
    "Overall rating: low if all domains >= 1 and total >= 8; high if total < 6; "
    "otherwise some concerns.")

rob = pd.read_csv(os.path.join(ROOT, "data", "processed", "risk_of_bias.csv"))
h("S5. Per-study risk-of-bias ratings", 1)
add_df(rob)

h("S6. Concentration-response slopes (per series)", 1)
S = pd.read_csv(os.path.join(ROOT, "results", "tables", "conc_response_slopes.csv"))
add_df(S, ["study_id", "honey", "organism", "n_conc", "spearman_rho", "p",
           "measure"])

h("S7. Natural-versus-artificial paired comparisons", 1)
P = pd.read_csv(os.path.join(ROOT, "results", "tables",
                             "natural_vs_artificial_pairs.csv"))
add_df(P)

h("S8. Catalase-treated versus untreated pairs", 1)
K = pd.read_csv(os.path.join(ROOT, "results", "tables", "catalase_pairs.csv"))
add_df(K)

h("S9. Chemistry-outcome correlations", 1)
C = pd.read_csv(os.path.join(ROOT, "results", "tables",
                             "chem_outcome_correlations.csv"))
add_df(C)

h("S10. Studies contributing extracted data", 1)
st = pd.read_csv(os.path.join(ROOT, "data", "processed", "studies.csv"))
add_df(st[st["extracted"]][["pmcid", "title", "journal", "year", "n_rows"]])

h("S11. Data limitations", 1)
doc.add_paragraph(
    "Concentration-series data: 8 of 103 extracted studies. Catalase-perturbed "
    "pairs: 13 observations, 2 studies. No extracted study reported H2O2 across "
    "a dilution ladder together with antibacterial endpoints; RQ3 is therefore "
    "supported only qualitatively. Water activity: 74 chemistry observations "
    "but few joinable to outcomes. Automated extraction was validated against "
    "manual curation on mechanistic tables; residual extraction error may "
    "remain, and every row carries table/row/column provenance.")

doc.save(os.path.join(OUT, "supplement.docx"))
print("wrote supplement.docx")

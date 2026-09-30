"""Generate the cover letter DOCX."""
import os

from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "manuscript")

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)

doc.add_paragraph("[Date]")
doc.add_paragraph()
doc.add_paragraph("Dear Editors,")
doc.add_paragraph()
doc.add_paragraph(
    "We submit our manuscript entitled \"Dilution-Dependent Antibacterial "
    "Mechanisms of Honey: A Systematic Review and Cross-Study Quantitative "
    "Evidence Synthesis\" for consideration in Antibiotics.")
doc.add_paragraph(
    "Honey's antibacterial activity is frequently assayed at high concentration, "
    "yet its component mechanisms respond differently to dilution: osmotic and "
    "acidity constraints weaken, whereas enzymatic hydrogen peroxide production "
    "can increase. Whether the dominant mechanism shifts with dilution has "
    "remained unclear because individual studies rarely measure all mechanisms "
    "across a concentration gradient. We synthesized this question across the "
    "published literature: a reproducible multi-database search (4,042 records, "
    "412 eligible studies) with structured extraction from 103 open-access full "
    "texts yielded 5,026 harmonized observations spanning antibacterial "
    "endpoints and physicochemical parameters, each with full provenance.")
doc.add_paragraph(
    "Key findings: extractable concentration-response series were uniformly "
    "positive (10/10); natural honey modestly but consistently outperformed "
    "sugar-equivalent controls; catalase-treated honey retained a median 92% of "
    "activity, indicating substantial non-peroxide mechanisms; methylglyoxal "
    "correlated inversely with MIC within the largest single-study block; and "
    "measured chemistry predicted antibacterial outcomes better than "
    "honey-type labels. We identify dilution-resolved peroxide measurement "
    "paired with antibacterial endpoints as a clear evidence gap.")
doc.add_paragraph(
    "The manuscript is accompanied by a full reproducible pipeline: raw search "
    "records, screening decisions, extraction code, the harmonized dataset with "
    "row-level provenance, a documented risk-of-bias framework, all analysis "
    "scripts, and figure-generation code.")
doc.add_paragraph(
    "This manuscript is original, has not been published previously, and is not "
    "under consideration elsewhere. All data derive from previously published "
    "public sources. The authors declare no conflicts of interest.")
doc.add_paragraph()
doc.add_paragraph("Sincerely,")
doc.add_paragraph("[Corresponding Author]")
doc.add_paragraph("[Affiliation]")

doc.save(os.path.join(OUT, "cover_letter.docx"))
print("wrote cover_letter.docx")

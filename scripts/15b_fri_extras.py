"""Generate FRI extras: graphical abstract (PNG+TIFF), highlights DOCX,
and cover letter DOCX.
"""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "results", "figures")
OUT = os.path.join(ROOT, "manuscript")
R = json.load(open(os.path.join(ROOT, "results", "models",
                                "analysis_results.json")))

# ================= GRAPHICAL ABSTRACT =================
fig, ax = plt.subplots(figsize=(7.0, 3.4), dpi=300)
ax.axis("off")
ax.set_xlim(0, 10)
ax.set_ylim(0, 5)

# concentration axis
ax.annotate("", xy=(9.4, 0.55), xytext=(0.5, 0.55),
            arrowprops=dict(arrowstyle="->", lw=1.4, color="black"))
ax.text(0.5, 0.15, "concentrated honey", fontsize=8, ha="left")
ax.text(9.4, 0.15, "diluted honey", fontsize=8, ha="right")
ax.text(4.95, 0.15, "decreasing concentration", fontsize=8, ha="center",
        style="italic")

# mechanism bands
bands = [
    ("osmotic / low water activity / acidity",
     (0.7, 4.4), (5.2, 4.4), "#33608a", "strong at high concentration,\nweakens on dilution"),
    ("hydrogen peroxide (glucose oxidase)",
     (3.2, 2.9), (9.0, 2.9), "#4a7c43", "enabled by dilution;\npeaks at intermediate strength"),
    ("non-peroxide factors (e.g. methylglyoxal)",
     (1.8, 1.4), (9.0, 1.4), "#8a5a33", "persists in some honeys;\nretained after catalase"),
]
for label, (x0, y0), (x1, y1), color, note in bands:
    ax.plot([x0, x1], [y0, y1], color=color, lw=6, solid_capstyle="round",
            alpha=0.75)
    ax.text(x0, y0 + 0.28, label, fontsize=8, color=color, weight="bold")
    ax.text(x1 + 0.1, y0, note, fontsize=6.5, color="#333333", va="center")

# evidence-gap marker
ax.plot([6.6], [2.9], marker="o", ms=6, mfc="none", mec="#4a7c43", mew=1.5)
ax.annotate("sparse direct dilution-resolved\nevidence (explicit gap)",
            xy=(6.6, 2.9), xytext=(7.6, 4.0), fontsize=6.5, color="#4a7c43",
            arrowprops=dict(arrowstyle="->", lw=0.7, color="#4a7c43"))

fig.savefig(os.path.join(FIG, "graphical_abstract.png"), dpi=300,
            bbox_inches="tight")
fig.savefig(os.path.join(FIG, "graphical_abstract.tiff"), dpi=500,
            bbox_inches="tight", format="tiff",
            pil_kwargs={"compression": "tiff_lzw"})
fig.savefig(os.path.join(FIG, "graphical_abstract.svg"),
            bbox_inches="tight")
plt.close(fig)
print("wrote graphical_abstract.png/.tiff")

# ================= HIGHLIGHTS =================
highlights = [
    "Honey antibacterial activity increased monotonically with concentration.",
    "Acidity and osmotic constraints weakened as honey was diluted.",
    "Natural honey exceeded sugar-equivalent controls in most matched pairs.",
    "Catalase-resistant activity suggests non-peroxide effects in tested honeys.",
    "Direct dilution-resolved peroxide measurements remain an evidence gap.",
]
for b in highlights:
    assert len(b) <= 85, (len(b), b)

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)
doc.add_heading("Highlights", 1)
doc.add_paragraph("Dilution-dependent antibacterial mechanisms of honey: a "
                  "systematic review and cross-study quantitative evidence "
                  "synthesis")
doc.add_paragraph()
for b in highlights:
    doc.add_paragraph(b, style="List Bullet")
doc.save(os.path.join(OUT, "highlights.docx"))
with open(os.path.join(OUT, "highlights.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join("- " + b for b in highlights) + "\n")
print("wrote highlights.docx/.txt;", [len(b) for b in highlights])

# ================= COVER LETTER =================
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)
doc.add_paragraph("[Date]")
doc.add_paragraph()
doc.add_paragraph("Dear Editor,")
doc.add_paragraph()
body = [
 "We submit our manuscript entitled \"Dilution-dependent antibacterial "
 "mechanisms of honey: a systematic review and cross-study quantitative "
 "evidence synthesis\" for consideration in Food Research International.",
 "Honey is widely studied as a natural antimicrobial ingredient, yet most "
 "in vitro evaluations are performed at a single concentration or over "
 "narrow ranges, and the mechanistic basis of activity is usually inferred "
 "rather than measured. Our review asks a question that sits squarely "
 "within the journal's scope at the interface of food functionality and "
 "food microbiology: does the mechanism behind honey's antibacterial "
 "activity change with dilution, rather than simply weakening?",
 "We systematically searched PubMed, Europe PMC and Crossref with "
 "programmatic citation mining and data-repository searches, screened "
 f"{R['selection']['eligible']} eligible primary studies, and extracted "
 f"{R['n_rows']:,} harmonized antibacterial and physicochemical "
 "observations from open-access full texts using a fully documented, "
 "reproducible pipeline. The synthesis shows that concentration-response "
 "is consistently positive while acidity weakens on dilution; that natural "
 "honey exceeds sugar-equivalent artificial controls in most matched "
 "pairs; that catalase-resistant residual activity is substantial where "
 "measured; and that within-study chemistry-outcome associations do not "
 "currently generalize across studies. We report the last point "
 "explicitly, together with a quantified description of which eligible "
 "studies could and could not be extracted, because we believe honest "
 "accounting of evidence strength is as valuable as the positive "
 "findings.",
 "The work is relevant to readers evaluating honey as a functional food "
 "ingredient or natural preservative, because assay concentration is "
 "shown to be mechanistically informative rather than a mere dosing "
 "detail. All data and code required to regenerate every reported number "
 "and figure are provided; no proprietary or fabricated data were used.",
 "The manuscript is original, has not been published previously, and is "
 "not under consideration elsewhere. All authors have approved the "
 "submission and declare no competing interests.",
 "We hope you will find the manuscript suitable for Food Research "
 "International and thank you for your consideration.",
]
for b in body:
    doc.add_paragraph(b)
    doc.add_paragraph()
doc.add_paragraph("Sincerely,")
doc.add_paragraph("[Corresponding Author]")
doc.add_paragraph("[Affiliation]")
doc.add_paragraph("[E-mail]")
doc.save(os.path.join(OUT, "cover_letter_fri.docx"))
print("wrote cover_letter_fri.docx")

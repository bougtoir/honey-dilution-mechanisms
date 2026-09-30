"""Assemble the Food Research International submission package.

Creates submission/fri_final/ with:
  - Fri_main_manuscript.docx   (standalone copy)
  - Fri_submission_package.zip (full package: manuscript files, figures,
    editable tables, data, code, environment files, audits, README)
"""
import json
import os
import shutil
import zipfile

import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FINAL = os.path.join(ROOT, "submission", "fri_final")
os.makedirs(FINAL, exist_ok=True)

R = json.load(open(os.path.join(ROOT, "results", "models",
                                "analysis_results.json")))

# ---------- standalone editable tables DOCX ----------
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)
doc.add_heading("Tables", 1)
doc.add_paragraph("Dilution-dependent antibacterial mechanisms of honey: a "
                  "systematic review and cross-study quantitative evidence "
                  "synthesis")
doc.add_paragraph()


def add_table(title, headers, rows, note=None):
    p = doc.add_paragraph()
    parts = title.split(". ", 1)
    r = p.add_run(parts[0] + ". ")
    r.bold = True
    p.add_run(parts[1] if len(parts) > 1 else "")
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, x in enumerate(headers):
        t.rows[0].cells[j].text = x
    for row in rows:
        c = t.add_row().cells
        for j, v in enumerate(row):
            c[j].text = str(v)
    if note:
        np_ = doc.add_paragraph()
        nr = np_.add_run("Note: " + note)
        nr.font.size = Pt(9)
    doc.add_paragraph()


ADIR = R["artificial_direction"]
A1, A2, A4 = R["A1"], R["A2"], R["A4"]
A3 = {c["param"] + "|" + c["outcome"]: c for c in R["A3_correlations"]}
mgo = A3.get("MGO|MIC", {})
phzone = A3.get("pH|zone", {})
PAIRS = pd.read_csv(os.path.join(ROOT, "results", "tables",
                               "natural_vs_artificial_pairs.csv"))

add_table("Table 1. Median MIC (% honey) by organism group, percentage-unit "
          "studies.",
          ["Organism group", "Median MIC (%)"],
          [(o, v) for o, v in R["A6"]["median_MIC_by_organism"].items()],
          note="MIC, minimum inhibitory concentration; MRSA/MSSA, "
               "methicillin-resistant/susceptible S. aureus. Values are "
               "medians of extracted MIC rows restricted to percentage-unit "
               "bases.")
def pstr(p):
    if p is None or (isinstance(p, float) and p != p):
        return "-"
    return "< 0.001" if p < 0.001 else (f"{p:.3f}" if p < 0.1 else f"{p:.2f}")


add_table("Table 2. Paired evidence for mechanistic contrasts.",
          ["Contrast", "n pairs/obs", "n studies", "Effect", "p value"],
          [
           ("Natural vs artificial honey (potency ratio)",
            str(ADIR["n_pairs"]), str(PAIRS["study_id"].nunique()),
            f"{ADIR['n_natural_stronger']}/{ADIR['n_pairs']} pairs > 1; "
            f"median {ADIR['median_ratio']:.2f}",
            pstr(ADIR["wilcoxon_p_logratio"])),
           ("Natural vs artificial honey, zone difference (mm)",
            str(A2["zone_pairs"]),
            str(PAIRS.loc[PAIRS["outcome"] == "zone", "study_id"].nunique()),
            f"median +{A2['zone_median_diff_mm']:.1f} mm",
            pstr(A2["zone_wilcoxon_p"])),
           ("Catalase residual activity", str(A4["n_pairs"]),
            str(len(A4["studies"])),
            f"median {A4['median_residual_frac']:.0%} retained", "-"),
           ("Activity vs concentration (series)",
            str(A1["n_activity_series"]),
            str(R["conc_response"]["n_studies"]),
            f"{A1['positive']}/{A1['n_activity_series']} positive; "
            f"median rho = {A1['median_rho']:.2f}", "-"),
           ("MGO vs MIC (within-study)", str(mgo.get("n", 0)),
            str(mgo.get("n_studies", 0)),
            f"rho = {mgo.get('spearman_rho', 0):.2f}", pstr(mgo.get("p"))),
           ("pH vs inhibition zone", str(phzone.get("n", 0)),
            str(phzone.get("n_studies", 0)),
            f"rho = {phzone.get('spearman_rho', 0):.2f}",
            pstr(phzone.get("p"))),
          ],
          note="Potency ratio > 1 indicates greater natural-honey activity "
               "(zone: natural/artificial; MIC: artificial/natural). MGO, "
               "methylglyoxal; MIC, minimum inhibitory concentration.")
tables_path = os.path.join(ROOT, "manuscript", "tables_fri.docx")
doc.save(tables_path)
print("wrote tables_fri.docx")

# ---------- package contents ----------
STAGE = os.path.join(FINAL, "pkg_tmp")
if os.path.exists(STAGE):
    shutil.rmtree(STAGE)

files = {
    "manuscript/Fri_main_manuscript.docx": "manuscript/manuscript_fri.docx",
    "manuscript/Fri_main_manuscript.pdf": "manuscript/manuscript_fri.pdf",
    "manuscript/Fri_cover_letter.docx": "manuscript/cover_letter_fri.docx",
    "manuscript/Fri_highlights.docx": "manuscript/highlights.docx",
    "manuscript/Fri_supplement.docx": "manuscript/supplement_fri.docx",
    "manuscript/Fri_tables.docx": "manuscript/tables_fri.docx",
    "data/extraction.csv": "data/processed/extraction.csv",
    "data/chemistry.csv": "data/processed/chemistry.csv",
    "data/studies.csv": "data/processed/studies.csv",
    "data/risk_of_bias.csv": "data/processed/risk_of_bias.csv",
    "data/data_dictionary.md": "data/processed/data_dictionary.md",
    "data/eligible.csv": "data/interim/eligible.csv",
    "data/screened.csv": "data/interim/screened.csv",
    "data/manual_extraction.csv": "data/interim/manual_extraction.csv",
    "data/search_log.csv": "logs/search_log.csv",
    "data/repo_search_log.csv": "logs/repo_search_log.csv",
    "data/prisma_counts.json": "logs/prisma_counts.json",
    "data/pmcid_resolution.json": "logs/pmcid_resolution.json",
    "references/references_fri.csv": "references/references_fri.csv",
    "references/references_fri_context.csv":
        "references/references_fri_context.csv",
    "results/analysis_results.json": "results/models/analysis_results.json",
    "results/analysis_summary.txt": "results/models/analysis_summary.txt",
    "results/analysis_summary_fri.txt":
        "results/models/analysis_summary_fri.txt",
    "README.md": "README.md",
    "environment.yml": "environment.yml",
    "requirements.txt": "requirements.txt",
    "Makefile": "Makefile",
}
for arc, rel in files.items():
    src = os.path.join(ROOT, rel)
    if os.path.exists(src):
        dst = os.path.join(STAGE, arc)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.copy2(src, dst)
    else:
        print("MISSING", rel)

# figures (PNG + TIFF + SVG)
for f in os.listdir(os.path.join(ROOT, "results", "figures")):
    if f.endswith((".png", ".tiff", ".svg")):
        dst = os.path.join(STAGE, "figures", f)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.copy2(os.path.join(ROOT, "results", "figures", f), dst)

# per-analysis result tables
for f in os.listdir(os.path.join(ROOT, "results", "tables")):
    dst = os.path.join(STAGE, "results", "tables", f)
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.copy2(os.path.join(ROOT, "results", "tables", f), dst)

# code
for f in sorted(os.listdir(os.path.join(ROOT, "scripts"))):
    if f.endswith(".py"):
        dst = os.path.join(STAGE, "code", f)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.copy2(os.path.join(ROOT, "scripts", f), dst)

# QC audits
for f in os.listdir(os.path.join(ROOT, "qc", "fri")):
    dst = os.path.join(STAGE, "qc", f)
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.copy2(os.path.join(ROOT, "qc", "fri", f), dst)

# ---------- zip ----------
zpath = os.path.join(FINAL, "Fri_submission_package.zip")
if os.path.exists(zpath):
    os.remove(zpath)
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dirpath, _, fns in os.walk(STAGE):
        for fn in fns:
            fp = os.path.join(dirpath, fn)
            z.write(fp, os.path.relpath(fp, STAGE))

# standalone manuscript copy
shutil.copy2(os.path.join(ROOT, "manuscript", "manuscript_fri.docx"),
             os.path.join(FINAL, "Fri_main_manuscript.docx"))

n = len(zipfile.ZipFile(zpath).namelist())
sz = os.path.getsize(zpath) / 1e6
msz = os.path.getsize(os.path.join(FINAL, "Fri_main_manuscript.docx")) / 1e6
print(f"wrote {zpath} ({n} files, {sz:.1f} MB)")
print(f"wrote Fri_main_manuscript.docx ({msz:.1f} MB)")
shutil.rmtree(STAGE)

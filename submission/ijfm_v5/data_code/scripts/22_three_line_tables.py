"""Formatting-only pass: convert all tables in the IJFM manuscript,
supplement and standalone tables docx to three-line style.

Removes the 'Light Grid Accent 1' style, all shading, all vertical and
internal horizontal borders; retains top rule, header-separator rule and
bottom rule. Repeats the header row across pages. Cell text is never
modified — verified by a pre/post text dump comparison.
"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def cell_texts(doc):
    out = []
    for t in doc.tables:
        out.append([[c.text for c in r.cells] for r in t.rows])
    return out


def clear_tbl_borders(tbl):
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    b = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{name}")
        if name in ("top", "bottom"):
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "8")          # 1 pt
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "none")
            e.set(qn("w:sz"), "0")
        b.append(e)
    tblPr.append(b)


def set_cell_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(el)
    b = OxmlElement("w:tcBorders")
    for name in ("top", "left", "right"):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        b.append(e)
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "8")
    e.set(qn("w:color"), "000000")
    b.append(e)
    tcPr.append(b)


def remove_shading(t):
    for el in list(t._tbl.iter(qn("w:shd"))):
        el.getparent().remove(el)


def three_line(doc):
    for t in doc.tables:
        tbl = t._tbl
        # drop the grid/accent style
        tblPr = tbl.tblPr
        for el in tblPr.findall(qn("w:tblStyle")):
            tblPr.remove(el)
        remove_shading(t)
        clear_tbl_borders(tbl)
        if len(t.rows) == 0:
            continue
        # header separator rule on every first-row cell + bold + repeat
        tr = t.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        if trPr.find(qn("w:tblHeader")) is None:
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        for c in t.rows[0].cells:
            set_cell_bottom_border(c)
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True


def inventory(doc, label):
    lines = []
    for i, t in enumerate(doc.tables):
        n_shd = len(list(t._tbl.iter(qn("w:shd"))))
        has_style = t._tbl.tblPr.find(qn("w:tblStyle")) is not None
        lines.append(
            f"| {label} T{i+1} | {len(t.rows)} x "
            f"{len(t.columns)} | style={has_style} | shading_cells={n_shd} |")
    return lines


JOBS = [
    ("manuscript/manuscript_ijfm.docx",
     "submission/ijfm_final/IJFM_main_manuscript_linenumbered_FINAL.docx",
     "main manuscript"),
    ("manuscript/supplement_ijfm.docx",
     "submission/ijfm_final/IJFM_supplement_FINAL.docx", "supplement"),
    ("manuscript/tables_ijfm.docx",
     "submission/ijfm_final/IJFM_tables_threeline.docx", "tables docx"),
]

baseline, ident = [], []
for src, dst, label in JOBS:
    sp = os.path.join(ROOT, src)
    if not os.path.exists(sp):
        print("MISSING", src)
        continue
    doc = Document(sp)
    baseline += inventory(doc, label)
    before = cell_texts(doc)
    three_line(doc)
    after = cell_texts(doc)
    ok = before == after
    ident.append(f"| {label} | tables={len(before)} | "
                 f"cell_text_identical={ok} |")
    assert ok, f"content changed in {label}"
    dp = os.path.join(ROOT, dst)
    doc.save(dp)
    print(f"wrote {dst}")

    # verify formatting result on the saved file
    d2 = Document(dp)
    v_shd = sum(len(list(t._tbl.iter(qn("w:shd"))))
                for t in d2.tables)
    v_style = sum(1 for t in d2.tables
                  if t._tbl.tblPr.find(qn("w:tblStyle")) is not None)
    ident.append(f"| {label} post-check | shading={v_shd} "
                 f"styles={v_style} |")

print("\nBASELINE\n" + "\n".join(baseline))
print("\nIDENTITY\n" + "\n".join(ident))
